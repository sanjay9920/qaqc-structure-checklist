import argparse
import json
import re
from pathlib import Path

import openpyxl


EXCLUDED_SHEETS = {"Index", "Index-1", "Checklist status", "Sheet2"}

DISPLAY_NAMES = {
    "LA": "Lightning Arrester",
    "ERT": "Earth Resistance Test",
    "IR": "Insulation Resistance",
    "IMP": "Module Performance Test (IMP)",
    "Pre-Fab": "Pre-Fabricated Building",
    "Inverter Pre-Commisioning": "Inverter Pre-Commissioning",
    "IDT-Aux Trafo Pre-Commisioning": "IDT / Auxiliary Transformer Pre-Commissioning",
    "HT Panel Pre-Commisioning": "HT Panel Pre-Commissioning",
    "Battery Charger Commisioning": "Battery Charger Commissioning",
    "Battery Commisioning": "Battery Commissioning",
    "Commisioning Report Scada": "SCADA Commissioning",
    "Commisioning Report ACDB": "ACDB Commissioning",
    "Commisioning Report UPS DB": "UPS DB Commissioning",
    "Survillance System": "Surveillance System",
    "Fire fightening": "Fire Fighting",
    "6 sqmm DC cable": "6 sq mm DC Cable",
    "LV IR": "LV Insulation Resistance",
}

FORMAT_OVERRIDES = {
    "6 sqmm DC cable": "CH-LP-EL-003/00",
    "Trench": "Not specified",
}

POINT_OVERRIDES = {
    "IMP": [f"IMP Reading {number:02d}" for number in range(1, 21)],
    "LV IR": [f"LV IR Reading R-{number}" for number in range(1, 9)],
    "ERT": [
        f"Earth Resistance Test Point {number:02d} (EP / W.G / W.O.G)"
        for number in range(1, 22)
    ],
}

SPECIAL_MEASUREMENT_FIELDS = {
    "ERT": [
        ("identification", "Identification", "", "C", True),
        ("wg", "W.G", "ohm", "F", False),
        ("wog", "W.O.G", "ohm", "G", False),
    ],
    "VOC Testing": [
        ("identification", "Identification", "", "C", True),
        ("module_wp", "Module Wp", "Wp", "E", False),
        ("polarity", "Polarity", "", "F", False),
        ("pe_voltage", "P-E", "V", "G", False),
        ("ne_floating_voltage", "N-E Floating Voltage", "V", "H", False),
        ("pe_floating_voltage", "P-E Floating Voltage", "V", "I", False),
        ("time", "Time", "", "J", False),
    ],
    "IR": [
        ("identification", "Identification", "", "C", True),
        ("cable_length", "Cable Length", "", "E", False),
        ("continuity", "Continuity", "", "F", False),
        ("r_e", "R-E", "ohm", "G", False),
        ("y_e", "Y-E", "ohm", "H", False),
        ("b_e", "B-E", "ohm", "I", False),
        ("r_y", "R-Y", "ohm", "J", False),
        ("y_b", "Y-B", "ohm", "K", False),
        ("b_r", "B-R", "ohm", "L", False),
        ("ir_positive", "IR (+ve)", "ohm", "M", False),
        ("ir_negative", "IR (-ve)", "ohm", "N", False),
    ],
    "IMP": [
        ("identification", "Identification", "", "C", True),
        ("imp", "IMP", "", "E", False),
        ("time", "Time", "", "G", False),
    ],
    "Inverter Pre-Commisioning": [
        ("inv_1", "INV-1", "", "G", False),
        ("inv_2", "INV-2", "", "H", False),
        ("inv_3", "INV-3", "", "I", False),
        ("inv_4", "INV-4", "", "J", False),
    ],
    "VOC": [
        ("identification", "Identification", "", "B", True),
        ("polarity", "Polarity", "", "D", False),
        ("p_e", "P-E", "V", "E", False),
        ("n_e", "N-E", "V", "F", False),
        ("p_n", "P-N", "V", "G", False),
    ],
    "LV IR": [
        ("lv1_r_e", "LV-1 R-E", "ohm", "C", False),
        ("lv1_y_e", "LV-1 Y-E", "ohm", "D", False),
        ("lv1_b_e", "LV-1 B-E", "ohm", "E", False),
        ("lv2_r_e", "LV-2 R-E", "ohm", "F", False),
        ("lv2_y_e", "LV-2 Y-E", "ohm", "G", False),
        ("lv2_b_e", "LV-2 B-E", "ohm", "H", False),
        ("lv3_r_e", "LV-3 R-E", "ohm", "I", False),
        ("lv3_y_e", "LV-3 Y-E", "ohm", "J", False),
        ("lv3_b_e", "LV-3 B-E", "ohm", "K", False),
        ("lv4_r_e", "LV-4 R-E", "ohm", "L", False),
        ("lv4_y_e", "LV-4 Y-E", "ohm", "M", False),
        ("lv4_b_e", "LV-4 B-E", "ohm", "N", False),
    ],
}

SPECIAL_POINT_ROWS = {
    "LV IR": list(range(12, 20)),
}

SPELLING_REPLACEMENTS = {
    "Balast": "Ballast",
    "Lighteng": "Lighting",
    "jumaer": "jumper",
    "Thtough": "Through",
    "Meggar": "Megger",
    "Verfiy": "Verify",
    "requirm": "requirem",
    "markin done": "marking done",
    "acssories": "accessories",
    "darwing": "drawing",
    "quantaity": "quantity",
    "switchyaed": "switchyard",
    "atation": "station",
    "tranformer": "transformer",
    "sysytem": "system",
    "volatge": "voltage",
    "Fittment": "Fitting",
    "Chequared": "Chequered",
    "availabe": "available",
    "connectded": "connected",
    "auxillary": "auxiliary",
    "asbuilt": "as-built",
}


def normalized_text(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def slug(value):
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def is_serial(value):
    if isinstance(value, (int, float)) and int(value) == value and value > 0:
        return True
    return bool(
        re.fullmatch(r"\d+(?:\.\d+)*\.?|\d+[a-z]?", str(value or "").strip(), re.I)
    )


def clean_point(value):
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    for old, new in SPELLING_REPLACEMENTS.items():
        text = text.replace(old, new)
    return text


def matching_column(ws, header_row, labels):
    for column in range(1, ws.max_column + 1):
        text = normalized_text(ws.cell(header_row, column).value)
        if any(label in text for label in labels):
            return column
    return None


def point_rows(ws):
    if ws.title in SPECIAL_POINT_ROWS:
        return SPECIAL_POINT_ROWS[ws.title]
    header_row, serial_column = find_header(ws)
    if not header_row:
        return []
    return [
        row
        for row in range(header_row + 1, ws.max_row + 1)
        if is_serial(ws.cell(row, serial_column).value)
    ]


def field_schema(key, label, unit="", default=""):
    hint = f"Enter {label}"
    if unit:
        hint += f" ({unit})"
    return {
        "key": key,
        "label": label,
        "unit": unit,
        "placeholder": hint,
        "default": clean_point(default) if default is not None else "",
    }


def point_metadata(ws, row):
    header_row, _serial_column = find_header(ws)
    check_type = ""
    acceptance_criteria = ""
    section = ""
    if header_row:
        check_column = matching_column(ws, header_row, ("type of check",))
        criteria_column = matching_column(
            ws, header_row, ("reference document", "acceptance criteria")
        )
        if check_column:
            check_type = clean_point(ws.cell(row, check_column).value)
        if criteria_column:
            acceptance_criteria = clean_point(ws.cell(row, criteria_column).value)

    if ws.title == "Commisioning Report Scada":
        for source_row in range(row, (header_row or row) - 1, -1):
            candidate = clean_point(ws.cell(source_row, 3).value)
            if candidate:
                section = candidate
                break

    if ws.title in SPECIAL_MEASUREMENT_FIELDS:
        fields = []
        for key, label, unit, column_letter, include_default in SPECIAL_MEASUREMENT_FIELDS[
            ws.title
        ]:
            default = ws[f"{column_letter}{row}"].value if include_default else ""
            fields.append(field_schema(key, label, unit, default))
    else:
        is_measurement = "measurement" in normalized_text(check_type)
        label = "Measured Value / Observation" if is_measurement else "Observation / Result"
        hint = (
            "Enter actual measured value with unit"
            if is_measurement
            else "Enter observation or result"
        )
        fields = [field_schema("observation", label)]
        fields[0]["placeholder"] = hint

    return {
        "check_type": check_type,
        "acceptance_criteria": acceptance_criteria,
        "section": section,
        "measurement_fields": fields,
        "value_enabled": True,
        "value_label": fields[0]["label"],
        "value_hint": fields[0]["placeholder"],
    }


def find_header(ws):
    for row in range(1, min(ws.max_row, 15) + 1):
        for column in range(1, ws.max_column + 1):
            text = normalized_text(ws.cell(row, column).value)
            if "sr no" in text or "sl no" in text or text == "s no":
                return row, column
    return None, None


def description_column(ws, header_row, serial_column):
    candidates = []
    for column in range(1, ws.max_column + 1):
        text = normalized_text(ws.cell(header_row, column).value)
        if any(
            label in text
            for label in ("description", "check points", "inspection", "identification")
        ):
            candidates.append(column)
    return max(candidates) if candidates else min(serial_column + 1, ws.max_column)


def observation_column(ws, header_row, description_col):
    for column in range(description_col + 1, ws.max_column + 1):
        text = normalized_text(ws.cell(header_row, column).value)
        if any(
            label in text
            for label in (
                "reference document",
                "acceptance criteria",
                "type of check",
                "observation",
                "result",
                "signature",
                "polarity",
                "remark",
                "cable length",
                "continuity",
                "time",
                "imp",
            )
        ):
            return column
    return ws.max_column + 1


def extract_points(ws):
    if ws.title in POINT_OVERRIDES:
        return POINT_OVERRIDES[ws.title]

    header_row, serial_column = find_header(ws)
    if not header_row:
        return []
    description_col = description_column(ws, header_row, serial_column)
    stop_col = observation_column(ws, header_row, description_col)
    points = []
    for row in range(header_row + 1, ws.max_row + 1):
        if not is_serial(ws.cell(row, serial_column).value):
            continue
        label = ""
        for column in range(description_col, stop_col):
            value = ws.cell(row, column).value
            if value is None or is_serial(value):
                continue
            candidate = clean_point(value)
            if candidate:
                label = candidate
                break
        if label:
            points.append(label)
    return points


def extract_format_number(ws):
    if ws.title in FORMAT_OVERRIDES:
        return FORMAT_OVERRIDES[ws.title]
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            match = re.search(r"TPS-LP-[A-Z]+-\d+(?:-\d+)?", cell.value)
            if match:
                return match.group(0)
    return "Not specified"


def build_catalog(workbook):
    templates = []
    for order, ws in enumerate(
        (sheet for sheet in workbook.worksheets if sheet.title not in EXCLUDED_SHEETS),
        start=1,
    ):
        name = DISPLAY_NAMES.get(ws.title, ws.title)
        template_id = slug(ws.title)
        points = extract_points(ws)
        if not points:
            raise ValueError(f"No checklist points found in sheet: {ws.title}")
        rows = point_rows(ws)
        if len(rows) < len(points):
            raise ValueError(
                f"Point row mapping failed in {ws.title}: {len(rows)} rows for {len(points)} points"
            )
        templates.append(
            {
                "template_id": template_id,
                "name": name,
                "source_sheet": ws.title,
                "format_no": extract_format_number(ws),
                "order": order * 10,
                "points": [
                    dict(
                        {
                        "item_id": f"point-{index:03d}",
                        "label": label,
                        "order": index * 10,
                        },
                        **point_metadata(ws, rows[index - 1]),
                    )
                    for index, label in enumerate(points, start=1)
                ],
            }
        )
    return {"version": 3, "template_count": len(templates), "templates": templates}


def main():
    parser = argparse.ArgumentParser(description="Extract equipment checklists from Excel.")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.input, data_only=True)
    catalog = build_catalog(workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"Saved {catalog['template_count']} checklist templates to {args.output}")


if __name__ == "__main__":
    main()

import copy
import sys
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

import checklist_app as app_module
from checklist_app.equipment_catalog import get_equipment_catalog, get_equipment_template
from checklist_app.equipment_services import (
    build_equipment_family_summaries,
    build_equipment_template_summaries,
    create_equipment_checklist,
    delete_equipment_checklist,
    get_equipment_history,
    get_next_equipment_record_number,
    list_equipment_checklists,
    update_equipment_final_remark,
    update_equipment_details,
    update_equipment_measurements,
    update_equipment_remark,
    update_equipment_status,
)
from checklist_app.exports import export_all_xlsx, export_equipment_csv
from checklist_app.services import (
    delete_checklist_item,
    delete_project,
    get_active_checklist_items,
)


class FakeSnapshot:
    def __init__(self, reference, data):
        self.reference = reference
        self.id = reference.id
        self._data = copy.deepcopy(data)
        self.exists = data is not None

    def to_dict(self):
        return copy.deepcopy(self._data)


class FakeDocument:
    def __init__(self, database, collection_name, document_id):
        self.database = database
        self.collection_name = collection_name
        self.id = document_id

    def get(self):
        return FakeSnapshot(
            self, self.database.data[self.collection_name].get(self.id)
        )

    def set(self, data):
        self.database.data[self.collection_name][self.id] = copy.deepcopy(data)

    def update(self, data):
        document = self.database.data[self.collection_name][self.id]
        for path, value in copy.deepcopy(data).items():
            parts = path.split(".")
            target = document
            for part in parts[:-1]:
                target = target.setdefault(part, {})
            target[parts[-1]] = value

    def delete(self):
        self.database.data[self.collection_name].pop(self.id, None)


class FakeQuery:
    def __init__(self, collection, field=None, value=None, maximum=None):
        self.collection = collection
        self.field = field
        self.value = value
        self.maximum = maximum

    def stream(self):
        rows = self.collection.stream()
        if self.field:
            rows = [
                row
                for row in rows
                if (row.to_dict() or {}).get(self.field) == self.value
            ]
        return rows[: self.maximum] if self.maximum else rows

    def order_by(self, *_args, **_kwargs):
        return self

    def limit(self, maximum):
        self.maximum = maximum
        return self


class FakeCollection:
    def __init__(self, database, name):
        self.database = database
        self.name = name
        database.data.setdefault(name, {})

    def document(self, document_id):
        return FakeDocument(self.database, self.name, document_id)

    def stream(self):
        return [
            FakeSnapshot(FakeDocument(self.database, self.name, key), value)
            for key, value in self.database.data[self.name].items()
        ]

    def add(self, data):
        reference = self.document(f"auto-{len(self.database.data[self.name]) + 1}")
        reference.set(data)
        return None, reference

    def where(self, filter=None, **_kwargs):
        return FakeQuery(
            self,
            getattr(filter, "field_path", None),
            getattr(filter, "value", None),
        )

    def order_by(self, *_args, **_kwargs):
        return FakeQuery(self)

    def limit(self, maximum):
        return FakeQuery(self, maximum=maximum)


class FakeBatch:
    def __init__(self):
        self.operations = []

    def delete(self, reference):
        self.operations.append(("delete", reference, None))

    def set(self, reference, data):
        self.operations.append(("set", reference, data))

    def commit(self):
        for action, reference, data in self.operations:
            if action == "delete":
                reference.delete()
            else:
                reference.set(data)
        self.operations = []


class FakeFirestore:
    def __init__(self):
        self.data = {}

    def collection(self, name):
        return FakeCollection(self, name)

    def batch(self):
        return FakeBatch()


def build_fixture():
    database = FakeFirestore()
    database.collection("checklist_items").document("dummy").set(
        {"label": "Dummy", "order": 10, "active": True}
    )
    user = {"uid": "admin", "email": "admin@example.com"}
    cable_laying = create_equipment_checklist(
        database,
        "100 MW AKOLA SITE",
        "1",
        "cable-laying",
        "01",
        user["email"],
        base_url="https://quality.example.com",
        vendor_name="Polycab",
        equipment_identification="SCB-1",
        specification="240 SQMM",
    )
    transformer = create_equipment_checklist(
        database,
        "100 MW AKOLA SITE",
        "1",
        "transformer-installation",
        "01",
        user["email"],
        base_url="https://quality.example.com",
    )
    cable_laying = update_equipment_status(
        database, cable_laying["equipment_id"], "point-001", "completed", user
    )
    cable_laying = update_equipment_remark(
        database,
        cable_laying["equipment_id"],
        "point-001",
        "Cable type verified",
        user,
    )
    cable_laying = update_equipment_final_remark(
        database,
        cable_laying["equipment_id"],
        "Ready for next inspection",
        user,
    )
    cable_laying = update_equipment_measurements(
        database,
        cable_laying["equipment_id"],
        "point-007",
        {"observation": "12 x cable OD"},
        user,
    )
    return database, user, cable_laying, transformer


def build_test_app(database, logged_in=False):
    admin = {
        "uid": "admin",
        "email": "admin@example.com",
        "is_admin": True,
        "all_projects": True,
        "projects": [],
    }
    app_module.get_db = lambda: database
    app_module.current_user = (lambda: admin) if logged_in else (lambda: None)
    return app_module.create_app()


class EquipmentWorkflowTests(unittest.TestCase):
    def test_equipment_family_summary_tracks_unique_units_and_all_work(self):
        records = [
            {
                "equipment_id": "record-1",
                "equipment_identification": "SCB-32",
                "template_id": "dc-cable-laying",
                "template_name": "DC Cable Laying",
                "block": "BLOCK-1",
                "counts": {"total": 10, "completed": 10, "pending": 0, "na": 0, "progress": 100},
            },
            {
                "equipment_id": "record-2",
                "equipment_identification": "SCB-32",
                "template_id": "dc-cable-termination",
                "template_name": "DC Cable Termination",
                "block": "BLOCK-1",
                "counts": {"total": 10, "completed": 3, "pending": 7, "na": 0, "progress": 30},
            },
            {
                "equipment_id": "record-3",
                "equipment_identification": "IDT-1",
                "template_id": "transformer-installation",
                "template_name": "Transformer Installation",
                "block": "BLOCK-2",
                "counts": {"total": 18, "completed": 0, "pending": 18, "na": 0, "progress": 0},
            },
            {
                "equipment_id": "record-4",
                "equipment_identification": "SCB-32",
                "template_id": "dc-cable-laying",
                "template_name": "DC Cable Laying",
                "block": "BLOCK-2",
                "counts": {"total": 10, "completed": 10, "pending": 0, "na": 0, "progress": 100},
            },
        ]

        summaries = build_equipment_family_summaries(records)
        scb = next(item for item in summaries if item["family"] == "SCB")
        idt = next(item for item in summaries if item["family"] == "IDT")

        self.assertEqual(scb["equipment_count"], 2)
        self.assertEqual(scb["completed_equipment"], 1)
        self.assertEqual(scb["in_progress_equipment"], 1)
        self.assertEqual(scb["pending_equipment"], 1)
        self.assertEqual(scb["checklist_type_count"], 2)
        self.assertEqual(scb["total_records"], 3)
        self.assertEqual(scb["completed_points"], 23)
        self.assertEqual(scb["pending_points"], 7)
        self.assertEqual(scb["equipment"][0]["label"], "SCB-32")
        self.assertEqual(scb["block_count"], 2)
        self.assertEqual(idt["equipment_count"], 1)
        self.assertEqual(idt["not_started_equipment"], 1)
        self.assertEqual(idt["block_count"], 1)

    def test_password_fields_have_show_hide_controls(self):
        database = FakeFirestore()
        client = build_test_app(database).test_client()

        page = client.get("/login")

        self.assertEqual(page.status_code, 200)
        self.assertIn(b'data-password-toggle="loginPassword"', page.data)
        self.assertIn(b'aria-label="Show password"', page.data)
        self.assertIn(b"password-toggle.js", page.data)
        service_worker = client.get("/static/service-worker.js")
        self.assertIn(b"quality-sims-v15", service_worker.data)
        self.assertIn(b"/static/js/password-toggle.js?v=2", service_worker.data)
        service_worker.close()

    def test_admin_can_recreate_email_from_legacy_removed_user(self):
        database = FakeFirestore()
        database.collection("checklist_items").document("dummy").set(
            {"label": "Dummy", "order": 10, "active": True}
        )
        database.collection("projects").document("PROJECT-1").set(
            {"project_id": "PROJECT-1", "display_name": "Project 1"}
        )
        client = build_test_app(database, logged_in=True).test_client()
        legacy_user = SimpleNamespace(uid="legacy-uid", disabled=True)
        new_user = SimpleNamespace(uid="new-uid")
        email_exists = app_module.firebase_auth.EmailAlreadyExistsError(
            "Email exists", None, None
        )

        with (
            patch.object(app_module, "initialize_firebase"),
            patch.object(
                app_module.firebase_auth,
                "create_user",
                side_effect=[email_exists, new_user],
            ) as create_user,
            patch.object(
                app_module.firebase_auth,
                "get_user_by_email",
                return_value=legacy_user,
            ),
            patch.object(app_module.firebase_auth, "delete_user") as delete_user,
            patch.object(app_module.firebase_auth, "set_custom_user_claims"),
        ):
            response = client.post(
                "/admin/users",
                data={
                    "name": "Recreated Worker",
                    "email": "worker@example.com",
                    "password": "secret123",
                    "role": "worker",
                    "projects": "PROJECT-1",
                },
            )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(create_user.call_count, 2)
        delete_user.assert_called_once_with("legacy-uid")

    def test_checklist_item_delete_is_permanent_and_cascades(self):
        database = FakeFirestore()
        database.collection("checklist_items").document("alignment").set(
            {"label": "Alignment checked", "order": 10, "active": False}
        )
        database.collection("structures").document("STR-0001").set(
            {
                "structure_id": "STR-0001",
                "checklist": {
                    "alignment": {"label": "Alignment checked", "status": "completed"},
                    "earthing": {"label": "Earthing completed", "status": "pending"},
                },
            }
        )
        database.collection("history").document("history-1").set(
            {"item_id": "alignment", "structure_id": "STR-0001"}
        )

        deleted = delete_checklist_item(database, "alignment", "admin@example.com")

        self.assertEqual(deleted["updated_structures"], 1)
        self.assertEqual(deleted["deleted_history"], 1)
        self.assertFalse(
            database.collection("checklist_items").document("alignment").get().exists
        )
        structure = (
            database.collection("structures").document("STR-0001").get().to_dict()
        )
        self.assertNotIn("alignment", structure["checklist"])
        self.assertIn("earthing", structure["checklist"])
        self.assertEqual(len(database.data["history"]), 0)

    def test_deleted_checklist_items_are_not_seeded_again(self):
        database = FakeFirestore()
        database.collection("checklist_items").document("only-item").set(
            {"label": "Only item", "order": 10, "active": True}
        )

        delete_checklist_item(database, "only-item", "admin@example.com")

        self.assertEqual(get_active_checklist_items(database), [])
        self.assertEqual(len(database.data["checklist_items"]), 0)

    def test_create_update_export_and_delete(self):
        database, user, record, _transformer = build_fixture()
        self.assertEqual(record["counts"]["completed"], 1)
        self.assertEqual(record["counts"]["pending"], 11)
        self.assertEqual(record["checklist"][0]["remark"], "Cable type verified")
        self.assertEqual(
            record["checklist"][6]["measurements"]["observation"],
            "12 x cable OD",
        )
        self.assertEqual(record["checklist"][6]["check_type"], "Measurement")
        self.assertEqual(record["equipment_identification"], "SCB-1")
        self.assertEqual(record["specification"], "240 SQMM")
        self.assertEqual(record["vendor_name"], "Polycab")
        self.assertEqual(len(get_equipment_history(database, record["equipment_id"])), 4)

        csv_text = export_equipment_csv(
            database, project="100 MW AKOLA SITE", block="1"
        )
        self.assertIn("Cable Laying", csv_text)
        self.assertIn("Cable type verified", csv_text)
        self.assertIn("12 x cable OD", csv_text)
        self.assertIn("Measured Value / Observation", csv_text)
        self.assertIn("SCB-1", csv_text)
        self.assertIn("240 SQMM", csv_text)
        self.assertIn("Polycab", csv_text)
        workbook = load_workbook(
            export_all_xlsx(database, project="100 MW AKOLA SITE", block="1")
        )
        self.assertEqual(
            workbook.sheetnames,
            ["Structures", "History", "Equipment Checklists", "Equipment History"],
        )
        self.assertIn(
            "Observations / Measurements",
            [cell.value for cell in workbook["Equipment Checklists"][1]],
        )
        self.assertIn(
            "Activity / Section",
            [cell.value for cell in workbook["Equipment Checklists"][1]],
        )

        deleted = delete_equipment_checklist(database, record["equipment_id"])
        self.assertEqual(deleted["equipment_id"], record["equipment_id"])
        self.assertEqual(
            len(list_equipment_checklists(database, project="100 MW AKOLA SITE")), 1
        )

    def test_public_page_api_and_qr_are_read_only(self):
        database, _user, record, _transformer = build_fixture()
        client = build_test_app(database).test_client()

        page = client.get(f"/equipment/{record['equipment_id']}")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"Cable Laying", page.data)
        self.assertIn(b"12 x cable OD", page.data)
        self.assertIn(b"SCB-1", page.data)
        self.assertIn(b"240 SQMM", page.data)
        self.assertIn(b"Polycab", page.data)
        self.assertIn(b"equipmentVendorHeader", page.data)
        self.assertIn(b"Vendor / Manufacturer", page.data)
        self.assertNotIn(b"Cable type verified", page.data)
        self.assertNotIn(b"Save Observation", page.data)

        payload = client.get(f"/api/equipment/{record['equipment_id']}").get_json()
        self.assertEqual(payload["equipment_identification"], "SCB-1")
        self.assertNotIn("remark", payload["checklist"][0])
        reading_point = next(
            item for item in payload["checklist"] if item["item_id"] == "point-007"
        )
        self.assertEqual(
            reading_point["measurements"]["observation"], "12 x cable OD"
        )
        self.assertNotIn("measurement_updated_by", reading_point)
        qr_response = client.get(f"/equipment/{record['equipment_id']}/qr.png")
        self.assertEqual(qr_response.status_code, 200)
        self.assertEqual(qr_response.mimetype, "image/png")

    def test_authorized_user_can_save_and_clear_reading(self):
        database, _user, record, _transformer = build_fixture()
        client = build_test_app(database, logged_in=True).test_client()
        endpoint = (
            f"/api/equipment/{record['equipment_id']}/items/point-007/measurements"
        )

        response = client.post(
            endpoint, json={"measurements": {"observation": "15 x cable OD"}}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json()["checklist"][6]["measurements"]["observation"],
            "15 x cable OD",
        )

        response = client.post(endpoint, json={"measurements": {"observation": ""}})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json()["checklist"][6]["measurements"]["observation"], ""
        )

        invalid = client.post(
            endpoint,
            json={"measurements": {"unknown_field": "Not allowed"}},
        )
        self.assertEqual(invalid.status_code, 400)

    def test_equipment_details_can_be_saved_and_cleared(self):
        database, user, record, _transformer = build_fixture()
        updated = update_equipment_details(
            database,
            record["equipment_id"],
            {
                "equipment_identification": "SCB-2",
                "specification": "400 SQMM",
                "vendor_name": "KEI",
            },
            user,
        )
        self.assertEqual(updated["equipment_identification"], "SCB-2")
        self.assertEqual(updated["specification"], "400 SQMM")
        self.assertEqual(updated["vendor_name"], "KEI")

        client = build_test_app(database, logged_in=True).test_client()
        partial = client.post(
            f"/api/equipment/{record['equipment_id']}/details",
            json={"vendor_name": "Polycab Limited"},
        )
        self.assertEqual(partial.status_code, 200)
        self.assertEqual(partial.get_json()["equipment_identification"], "SCB-2")
        self.assertEqual(partial.get_json()["specification"], "400 SQMM")
        self.assertEqual(partial.get_json()["vendor_name"], "Polycab Limited")

        response = client.post(
            f"/api/equipment/{record['equipment_id']}/details",
            json={
                "equipment_identification": "",
                "specification": "",
                "vendor_name": "",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["vendor_name"], "")
        history = get_equipment_history(database, record["equipment_id"])
        self.assertEqual(history[0]["change_type"], "details")

        too_long = client.post(
            f"/api/equipment/{record['equipment_id']}/details",
            json={"vendor_name": "X" * 121},
        )
        self.assertEqual(too_long.status_code, 400)

    def test_every_excel_checklist_has_correct_observation_schema(self):
        catalog = get_equipment_catalog()
        self.assertEqual(len(catalog), 45)
        self.assertEqual(sum(len(item["points"]) for item in catalog), 586)
        names = {item["name"] for item in catalog}
        self.assertTrue(
            {
                "AC Cable Laying",
                "DC Cable Laying",
                "AC Cable Termination",
                "DC Cable Termination",
            }.issubset(names)
        )
        self.assertTrue(
            all(point.get("measurement_fields") for item in catalog for point in item["points"])
        )
        expected_field_counts = {
            "ert": 3,
            "voc-testing": 7,
            "ir": 11,
            "imp": 3,
            "inverter-pre-commisioning": 4,
            "voc": 5,
            "lv-ir": 12,
        }
        for template_id, field_count in expected_field_counts.items():
            template = get_equipment_template(template_id)
            self.assertEqual(len(template["points"][0]["measurement_fields"]), field_count)

    def test_every_excel_checklist_can_save_all_observation_fields(self):
        database = FakeFirestore()
        user = {"email": "qa@example.com", "uid": "qa-user"}

        for template in get_equipment_catalog():
            record = create_equipment_checklist(
                database,
                "100 MW AKOLA SITE",
                "1",
                template["template_id"],
                "01",
                user["email"],
                base_url="https://quality.example.com",
            )
            first_point = record["checklist"][0]
            values = {
                field["key"]: f"Test {field['label']}"
                for field in first_point["measurement_fields"]
            }
            updated = update_equipment_measurements(
                database,
                record["equipment_id"],
                first_point["item_id"],
                values,
                user,
            )
            self.assertEqual(updated["checklist"][0]["measurements"], values)

    def test_block_dashboard_contains_catalog_and_tracking(self):
        database, _user, record, _transformer = build_fixture()
        client = build_test_app(database, logged_in=True).test_client()
        url = "/admin?project=100-MW-AKOLA-SITE&block=BLOCK-1"
        page = client.get(url)
        self.assertEqual(page.status_code, 200)
        for expected in [
            b"45 types",
            b"AC Cable Laying",
            b"DC Cable Laying",
            b"AC Cable Termination",
            b"DC Cable Termination",
            b"12 points",
            b"Completed points",
            b"SCB-1",
            b"240 SQMM",
            b"Checklist dashboards",
            b"45 dashboards",
            b"Open Dashboard",
            b"in progress",
            b"not started",
            b"ID missing",
            b"Search checklist, equipment or vendor",
            b"Equipment and circuit register",
            b"Created checklist type status",
        ]:
            self.assertIn(expected, page.data)

        payload = client.get(
            "/admin/api/structures?project=100-MW-AKOLA-SITE&block=BLOCK-1"
        ).get_json()
        self.assertEqual(payload["equipment_summary"]["total_records"], 2)
        self.assertEqual(payload["equipment_records"][0]["equipment_id"], record["equipment_id"])
        self.assertEqual(len(payload["equipment_template_summaries"]), 45)
        cable_summary = next(
            item
            for item in payload["equipment_template_summaries"]
            if item["template_id"] == "cable-laying"
        )
        self.assertEqual(cable_summary["in_progress_records"], 1)
        self.assertEqual(cable_summary["not_started_records"], 0)
        self.assertEqual(cable_summary["identified_records"], 1)
        self.assertEqual(cable_summary["missing_identity_records"], 0)
        self.assertEqual(cable_summary["equipment_labels"], ["SCB-1"])
        self.assertEqual(cable_summary["vendor_names"], ["Polycab"])
        self.assertEqual(cable_summary["specifications"], ["240 SQMM"])
        self.assertNotIn(b"Created equipment checklists", page.data)
        self.assertNotIn(b"Equipment History CSV", page.data)

    def test_project_dashboard_combines_every_blocks_structure_and_equipment_work(self):
        database, _user, _record, _transformer = build_fixture()
        client = build_test_app(database, logged_in=True).test_client()
        url = "/admin?project=100-MW-AKOLA-SITE"

        page = client.get(url)
        self.assertEqual(page.status_code, 200)
        for expected in [
            b"Project QA/QC control center",
            b"All checklist and work status",
            b"Every block: checklist and work details",
            b"Block 1",
            b"Equipment checklists",
            b"Combined QA/QC work",
            b"Open Block Dashboard",
            b"Equipment and circuit register",
            b"Created checklist type status",
        ]:
            self.assertIn(expected, page.data)
        self.assertNotIn(b"Open block", page.data)
        self.assertNotIn(b"Block detail", page.data)

        payload = client.get(
            "/admin/api/structures?project=100-MW-AKOLA-SITE"
        ).get_json()
        summary = payload["project_summary"]
        self.assertEqual(summary["equipment_total_records"], 2)
        self.assertEqual(summary["equipment_pending_records"], 2)
        self.assertEqual(summary["equipment_total_points"], 30)
        self.assertEqual(summary["equipment_completed_points"], 1)
        self.assertEqual(summary["work_total_points"], 30)
        self.assertEqual(summary["work_completed_points"], 1)
        block = next(item for item in summary["blocks"] if item["block"] == "BLOCK-1")
        self.assertEqual(block["equipment_total_records"], 2)
        self.assertEqual(block["equipment_pending_records"], 2)
        self.assertEqual(block["work_progress"], 3)

    def test_admin_project_delete_cascades_related_data(self):
        database = FakeFirestore()
        project_id = "150-MW-AKOLA-SITE"
        structure_id = f"{project_id}-BLOCK-1-STR-01"
        database.collection("projects").document(project_id).set(
            {"project_id": project_id}
        )
        database.collection("structures").document(structure_id).set(
            {"project": project_id, "block": "BLOCK-1"}
        )
        database.collection("equipment_checklists").document("equipment-1").set(
            {"project": project_id, "block": "BLOCK-1"}
        )
        database.collection("history").document("history-1").set(
            {"project": project_id, "structure_id": structure_id}
        )
        database.collection("equipment_history").document("equipment-history-1").set(
            {"project": project_id, "equipment_id": "equipment-1"}
        )

        deleted = delete_project(database, project_id, cascade=True)

        self.assertEqual(deleted["structures_deleted"], 1)
        self.assertEqual(deleted["equipment_deleted"], 1)
        self.assertFalse(database.collection("projects").document(project_id).get().exists)
        self.assertFalse(
            database.collection("structures").document(structure_id).get().exists
        )
        self.assertFalse(
            database.collection("equipment_checklists")
            .document("equipment-1")
            .get()
            .exists
        )
        self.assertFalse(
            database.collection("history").document("history-1").get().exists
        )
        self.assertFalse(
            database.collection("equipment_history")
            .document("equipment-history-1")
            .get()
            .exists
        )

    def test_checklist_dashboard_supports_multiple_equipment_records(self):
        database, _user, record, _transformer = build_fixture()
        client = build_test_app(database, logged_in=True).test_client()
        dashboard_url = (
            "/admin/equipment-dashboard/cable-laying"
            "?project=100-MW-AKOLA-SITE&block=BLOCK-1"
        )

        page = client.get(dashboard_url)
        self.assertEqual(page.status_code, 200)
        for expected in [
            b"Cable Laying",
            b"SCB-1",
            b"240 SQMM",
            b"Polycab",
            b'value="02"',
            b"Vendor / Manufacturer",
            b"In progress",
            b"Not started",
            b"Unique equipment",
            b"Search equipment, specification or vendor",
        ]:
            self.assertIn(expected, page.data)

        created = client.post(
            "/admin/equipment",
            data={
                "project": "100-MW-AKOLA-SITE",
                "block": "BLOCK-1",
                "template_id": "cable-laying",
                "record_number": "",
                "equipment_identification": "SCB-2",
                "specification": "400 SQMM",
                "vendor_name": "KEI",
            },
            follow_redirects=False,
        )
        self.assertEqual(created.status_code, 302)
        self.assertIn("CABLE-LAYING-02", created.headers["Location"])

        payload = client.get(
            "/admin/api/equipment-dashboard/cable-laying"
            "?project=100-MW-AKOLA-SITE&block=BLOCK-1"
        ).get_json()
        self.assertEqual(payload["summary"]["total_records"], 2)
        self.assertEqual(payload["summary"]["pending_records"], 2)
        self.assertEqual(payload["next_record_number"], "03")
        scb_two = next(
            item
            for item in payload["records"]
            if item["equipment_identification"] == "SCB-2"
        )
        self.assertEqual(scb_two["record_number"], "02")
        self.assertEqual(scb_two["specification"], "400 SQMM")
        self.assertEqual(scb_two["vendor_name"], "KEI")

        summaries = build_equipment_template_summaries(
            get_equipment_catalog(), payload["records"]
        )
        cable_summary = next(
            item for item in summaries if item["template_id"] == "cable-laying"
        )
        self.assertEqual(cable_summary["total_records"], 2)
        self.assertEqual(
            get_next_equipment_record_number(payload["records"]), "03"
        )

    def test_security_headers_safe_redirect_and_missing_structure_qr(self):
        database, _user, _record, _transformer = build_fixture()
        public_client = build_test_app(database).test_client()
        response = public_client.get("/login")
        self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
        self.assertEqual(response.headers["X-Frame-Options"], "DENY")
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.assertEqual(
            public_client.get("/structure/DOES-NOT-EXIST/qr.png").status_code,
            404,
        )

        logged_in_client = build_test_app(database, logged_in=True).test_client()
        redirected = logged_in_client.get(
            "/login?next=https://example.com/phishing", follow_redirects=False
        )
        self.assertEqual(redirected.status_code, 302)
        self.assertEqual(redirected.headers["Location"], "/admin")

    def test_exports_neutralize_spreadsheet_formulas(self):
        database, user, record, _transformer = build_fixture()
        update_equipment_details(
            database,
            record["equipment_id"],
            {
                "equipment_identification": "=HYPERLINK(\"https://example.com\")",
                "specification": "240 SQMM",
                "vendor_name": "Polycab",
            },
            user,
        )
        csv_text = export_equipment_csv(database, project="100 MW AKOLA SITE")
        self.assertIn("'=HYPERLINK", csv_text)


if __name__ == "__main__":
    if "--preview" in sys.argv:
        preview_db, preview_user, preview_record, _preview_transformer = build_fixture()
        preview_ir = create_equipment_checklist(
            preview_db,
            "100 MW AKOLA SITE",
            "1",
            "ir",
            "01",
            preview_user["email"],
            base_url="https://quality.example.com",
        )
        preview_ir = update_equipment_measurements(
            preview_db,
            preview_ir["equipment_id"],
            "point-001",
            {
                "identification": "SCB-1",
                "cable_length": "250 m",
                "continuity": "OK",
                "r_e": "850",
                "y_e": "870",
                "b_e": "860",
            },
            preview_user,
        )
        preview_app = build_test_app(preview_db, logged_in=True)
        print(f"Preview equipment ID: {preview_record['equipment_id']}")
        print(f"Preview IR ID: {preview_ir['equipment_id']}")
        preview_app.run(host="127.0.0.1", port=5001, debug=False)
    else:
        unittest.main()
